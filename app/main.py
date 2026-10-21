import os
import re
import uuid
import sqlite3
from html import escape
from pathlib import Path
from urllib.parse import quote

from fastapi import (
    FastAPI,
    Request,
    Form,
    UploadFile,
    File,
    HTTPException,
)
from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    FileResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import load_workbook

from .db import connect, init_db
from .security import verify_admin, make_code, load_env_file


# =========================================================
# CẤU HÌNH
# =========================================================

BASE_DIR = Path(__file__).resolve().parent.parent

load_env_file(BASE_DIR / ".env")

TEMPLATES_DIR = BASE_DIR / "app" / "templates"
STATIC_DIR = BASE_DIR / "app" / "static"

UPLOAD_DIR = Path(
    os.getenv("UPLOAD_DIR", str(BASE_DIR / "uploads"))
).expanduser()

if not UPLOAD_DIR.is_absolute():
    UPLOAD_DIR = (BASE_DIR / UPLOAD_DIR).resolve()

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="Hệ thống điểm danh hội nghị"
)


# =========================================================
# BẢO MẬT SESSION
# =========================================================

APP_SECRET = os.getenv("APP_SECRET", "").strip()

if not APP_SECRET:
    # Chỉ dùng để chạy thử trên máy cá nhân.
    # Khi triển khai Render, phải cấu hình APP_SECRET.
    APP_SECRET = "development-only-secret-change-me"

app.add_middleware(
    SessionMiddleware,
    secret_key=APP_SECRET,
    same_site="lax",
    https_only=(
        os.getenv("COOKIE_HTTPS_ONLY", "false").lower()
        == "true"
    ),
    max_age=60 * 60 * 8,
)


# =========================================================
# TEMPLATE VÀ STATIC
# =========================================================

templates = Jinja2Templates(
    directory=str(TEMPLATES_DIR)
)

if STATIC_DIR.exists():
    app.mount(
        "/static",
        StaticFiles(directory=str(STATIC_DIR)),
        name="static",
    )


# Khởi tạo database khi ứng dụng chạy.
init_db()


# =========================================================
# HÀM DÙNG CHUNG
# =========================================================

def is_admin(request: Request) -> bool:
    """Kiểm tra phiên đăng nhập quản trị viên."""

    return request.session.get("role") == "admin"


def normalized(value) -> str:
    """Chuẩn hóa khoảng trắng trong dữ liệu đầu vào."""

    return re.sub(
        r"\s+",
        " ",
        str(value or "")
    ).strip()


def render(
    request: Request,
    template: str,
    context: dict,
    status_code: int = 200,
):
    """Render template và truyền trạng thái đăng nhập admin."""

    base_context = {
        "request": request,
        "is_admin": is_admin(request),
    }

    base_context.update(context)

    return templates.TemplateResponse(
        request=request,
        name=template,
        context=base_context,
        status_code=status_code,
    )


def admin_login_redirect():
    """Chuyển người chưa đăng nhập đến trang admin login."""

    return RedirectResponse(
        url="/admin/login",
        status_code=303,
    )


def admin_required(request: Request):
    """
    Bảo vệ các chức năng quản trị.
    Gọi hàm này đầu mỗi route chỉ dành cho admin.
    """

    if not is_admin(request):
        raise HTTPException(
            status_code=403,
            detail="Bạn không có quyền thực hiện thao tác này.",
        )


# =========================================================
# TRANG CHỦ CÔNG KHAI
# =========================================================

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    """
    Trang chủ công khai.

    Không yêu cầu đại biểu đăng nhập.
    Mọi người đều có thể xem danh sách văn bản.
    """

    with connect() as conn:
        documents = conn.execute("""
            SELECT *
            FROM documents
            ORDER BY uploaded_at DESC, id DESC
        """).fetchall()

    return render(
        request,
        "participant.html",
        {
            "page_title": "Văn bản hội nghị",
            "person": None,
            "documents": documents,
        },
    )


# =========================================================
# URL ĐĂNG NHẬP CŨ CỦA ĐẠI BIỂU
# =========================================================

@app.get("/login")
def old_login_page():
    """
    Chuyển URL đăng nhập đại biểu cũ sang trang admin login.
    """

    return RedirectResponse(
        "/admin/login",
        status_code=303,
    )


@app.post("/login")
def old_login_submit():
    """
    Không còn đăng nhập bằng mã đại biểu và họ tên.
    """

    return RedirectResponse(
        "/admin/login",
        status_code=303,
    )


# =========================================================
# ĐĂNG NHẬP QUẢN TRỊ VIÊN
# =========================================================

@app.get(
    "/admin/login",
    response_class=HTMLResponse,
)
def admin_login_page(request: Request):
    """Hiển thị trang đăng nhập quản trị."""

    if is_admin(request):
        return RedirectResponse(
            "/admin",
            status_code=303,
        )

    return render(
        request,
        "admin_login.html",
        {
            "page_title": "Đăng nhập quản trị",
            "error": None,
        },
    )


@app.post(
    "/admin/login",
    response_class=HTMLResponse,
)
def admin_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
):
    """Xác thực tài khoản admin."""

    username = normalized(username)

    if not verify_admin(username, password):
        return render(
            request,
            "admin_login.html",
            {
                "page_title": "Đăng nhập quản trị",
                "error": (
                    "Tên đăng nhập hoặc mật khẩu không đúng. "
                    "Vui lòng kiểm tra cấu hình máy chủ."
                ),
            },
            status_code=400,
        )

    # Xóa session cũ trước khi tạo phiên admin.
    request.session.clear()
    request.session["role"] = "admin"

    return RedirectResponse(
        "/admin",
        status_code=303,
    )


# =========================================================
# ĐĂNG XUẤT
# =========================================================

@app.post("/logout")
def logout(request: Request):
    """Đăng xuất và xóa session hiện tại."""

    request.session.clear()

    return RedirectResponse(
        "/",
        status_code=303,
    )


# =========================================================
# DANH SÁCH VĂN BẢN CÔNG KHAI
# =========================================================

@app.get(
    "/documents",
    response_class=HTMLResponse,
)
def documents_page(request: Request):
    """
    Giữ lại URL /documents cũ.

    Trang văn bản được công khai cho mọi người.
    """

    with connect() as conn:
        documents = conn.execute("""
            SELECT *
            FROM documents
            ORDER BY uploaded_at DESC, id DESC
        """).fetchall()

    return render(
        request,
        "participant.html",
        {
            "page_title": "Văn bản hội nghị",
            "person": None,
            "documents": documents,
        },
    )


# =========================================================
# TRANG QUẢN TRỊ
# =========================================================

@app.get(
    "/admin",
    response_class=HTMLResponse,
)
def admin_page(
    request: Request,
    q: str = "",
    status: str = "all",
    message: str = "",
    error: str = "",
):
    """Trang quản lý đại biểu, điểm danh và văn bản."""

    if not is_admin(request):
        return admin_login_redirect()

    q = normalized(q)

    clauses = []
    params = []

    # Tìm kiếm đại biểu.
    if q:
        keyword = f"%{q}%"

        clauses.append("""
            (
                full_name LIKE ?
                OR code LIKE ?
                OR organization LIKE ?
                OR role LIKE ?
                OR position LIKE ?
            )
        """)

        params.extend([
            keyword,
            keyword,
            keyword,
            keyword,
            keyword,
        ])

    # Lọc theo trạng thái điểm danh.
    if status == "present":
        clauses.append("checked_in = 1")

    elif status == "absent":
        clauses.append("checked_in = 0")

    sql = "SELECT * FROM participants"

    if clauses:
        sql += " WHERE " + " AND ".join(clauses)

    sql += """
        ORDER BY
            CASE
                WHEN role LIKE '%Chủ trì%' THEN 0
                ELSE 1
            END,
            full_name COLLATE NOCASE
    """

    with connect() as conn:
        people = conn.execute(
            sql,
            params,
        ).fetchall()

        total = conn.execute("""
            SELECT COUNT(*) AS n
            FROM participants
        """).fetchone()["n"]

        present = conn.execute("""
            SELECT COUNT(*) AS n
            FROM participants
            WHERE checked_in = 1
        """).fetchone()["n"]

        documents = conn.execute("""
            SELECT *
            FROM documents
            ORDER BY uploaded_at DESC, id DESC
        """).fetchall()

    return render(
        request,
        "admin.html",
        {
            "page_title": "Quản trị hội nghị",
            "people": people,
            "documents": documents,
            "total": total,
            "present": present,
            "absent": total - present,
            "q": q,
            "status": status,
            "message": message,
            "error": error,
        },
    )


# =========================================================
# IN DANH SÁCH ĐIỂM DANH RIÊNG
# =========================================================

@app.get("/admin/print-attendance", response_class=HTMLResponse)
def print_attendance(request: Request):
    """Hiển thị trang in chỉ gồm danh sách đại biểu và trạng thái điểm danh."""

    if not is_admin(request):
        return admin_login_redirect()

    with connect() as conn:
        people = conn.execute("""
            SELECT full_name, position, checked_in
            FROM participants
            ORDER BY full_name COLLATE NOCASE
        """).fetchall()

    total = len(people)
    present = sum(1 for person in people if person["checked_in"])
    absent = total - present

    rows = []
    for index, person in enumerate(people, start=1):
        name = escape(str(person["full_name"] or ""))
        position = escape(str(person["position"] or "—"))
        if person["checked_in"]:
            status_text = "CÓ MẶT"
            status_class = "present"
        else:
            status_text = "VẮNG MẶT"
            status_class = "absent"

        rows.append(
            f"<tr><td class='center'>{index}</td>"
            f"<td>{name}</td><td>{position}</td>"
            f"<td class='center status {status_class}'>{status_text}</td></tr>"
        )

    table_rows = "".join(rows) or (
        "<tr><td colspan='4' class='empty'>Chưa có đại biểu trong danh sách.</td></tr>"
    )

    page = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Danh sách điểm danh</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ font-family: Arial, Helvetica, sans-serif; color: #111827; margin: 0; padding: 28px; }}
  .toolbar {{ display: flex; gap: 10px; justify-content: flex-end; margin: 0 auto 18px; max-width: 1000px; }}
  .toolbar a, .toolbar button {{ border: 0; border-radius: 6px; padding: 10px 16px; font-size: 14px; cursor: pointer; text-decoration: none; }}
  .print {{ background: #166534; color: white; }} .back {{ background: #e5e7eb; color: #111827; }}
  .sheet {{ max-width: 1000px; margin: auto; }}
  h1 {{ text-align: center; font-size: 22px; margin: 0 0 8px; text-transform: uppercase; }}
  .subtitle {{ text-align: center; margin: 0 0 22px; font-size: 13px; color: #4b5563; }}
  .summary {{ display: flex; gap: 24px; justify-content: center; margin-bottom: 18px; font-size: 14px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  th, td {{ border: 1px solid #6b7280; padding: 9px 10px; text-align: left; }}
  th {{ background: #f3f4f6; text-align: center; }}
  .center {{ text-align: center; }} .status {{ font-weight: 700; }}
  .present {{ color: #166534; }} .absent {{ color: #b91c1c; }}
  .empty {{ text-align: center; padding: 20px; }}
  @media print {{
    @page {{ size: A4 portrait; margin: 12mm; }}
    body {{ padding: 0; }} .toolbar {{ display: none !important; }}
    .sheet {{ max-width: none; }}
    thead {{ display: table-header-group; }} tr {{ break-inside: avoid; }}
    th {{ background: #eee !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
  }}
</style>
</head>
<body>
  <div class="toolbar">
    <a class="back" href="/admin">Quay lại quản trị</a>
    <button class="print" onclick="window.print()">In danh sách</button>
  </div>
  <main class="sheet">
    <h1>DANH SÁCH ĐIỂM DANH HỘI NGHỊ</h1>
    <p class="subtitle">Bảng tổng hợp tình trạng có mặt và vắng mặt</p>
    <div class="summary">
      <span>Tổng số: <strong>{total}</strong></span>
      <span>Có mặt: <strong>{present}</strong></span>
      <span>Vắng mặt: <strong>{absent}</strong></span>
    </div>
    <table>
      <thead><tr><th style="width: 55px">STT</th><th>Họ và tên</th><th style="width: 28%">Chức vụ</th><th style="width: 150px">Tình trạng</th></tr></thead>
      <tbody>{table_rows}</tbody>
    </table>
  </main>
</body>
</html>"""
    return HTMLResponse(content=page)


# =========================================================
# NHẬP DANH SÁCH ĐẠI BIỂU TỪ EXCEL
# =========================================================


@app.post("/admin/import-excel")
async def import_excel(request: Request, file: UploadFile = File(...)):
    """
    Nhập danh sách đại biểu từ Excel.

    Mỗi lần nhập thành công, danh sách participants hiện tại sẽ được thay
    thế hoàn toàn bằng danh sách trong file Excel. Nếu quá trình ghi database
    thất bại, giao dịch sẽ rollback để giữ nguyên danh sách cũ.
    """
    if not is_admin(request):
        return admin_login_redirect()

    if (
        not file.filename
        or Path(file.filename).suffix.lower() not in {".xlsx", ".xlsm"}
    ):
        return RedirectResponse(
            "/admin?error=" + quote("Vui lòng tải lên file Excel .xlsx hoặc .xlsm."),
            status_code=303,
        )

    tmp_path = UPLOAD_DIR / f"import_{uuid.uuid4().hex}.xlsx"
    wb = None

    try:
        content = await file.read()

        if not content:
            raise ValueError("File Excel đang trống.")

        if len(content) > 10 * 1024 * 1024:
            raise ValueError("File Excel vượt quá giới hạn 10 MB.")

        tmp_path.write_bytes(content)

        wb = load_workbook(tmp_path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)

        # Đọc tiêu đề cột.
        headers = [
            normalized(value).casefold()
            for value in next(rows, ())
        ]

        aliases = {
            "stt": {"stt", "số thứ tự", "so thu tu"},
            "name": {
                "tên", "họ tên", "họ và tên", "tên đại biểu",
                "ten", "ho ten", "ho va ten", "ten dai bieu",
                "name", "full name",
            },
            "position": {
                "chức vụ", "chuc vu", "position", "title",
            },
        }

        indexes = {}
        for key, names in aliases.items():
            for index, header in enumerate(headers):
                if header in names:
                    indexes[key] = index
                    break

        missing = [
            label
            for key, label in (
                ("stt", "STT"),
                ("name", "Tên"),
                ("position", "Chức vụ"),
            )
            if key not in indexes
        ]

        if missing:
            raise ValueError(
                "Thiếu cột: "
                + ", ".join(missing)
                + ". File Excel cần có các cột STT, Tên, Chức vụ."
            )

        # Đọc và kiểm tra toàn bộ dữ liệu trước khi đụng đến danh sách cũ.
        records = []

        for row_number, row in enumerate(rows, start=2):
            def cell(key):
                index = indexes[key]
                if index >= len(row) or row[index] is None:
                    return ""
                return normalized(row[index])

            name = cell("name")
            position = cell("position")
            raw_stt = cell("stt")

            # Bỏ qua dòng trống hoàn toàn.
            if not name and not position and not raw_stt:
                continue

            if not name:
                raise ValueError(
                    f"Dòng {row_number} có dữ liệu nhưng thiếu Tên."
                )

            try:
                stt = float(raw_stt.replace(",", ".")) if raw_stt else float("inf")
            except ValueError:
                stt = float("inf")

            records.append({
                "stt": stt,
                "name": name,
                "position": position,
            })

        if not records:
            raise ValueError("File Excel không có đại biểu hợp lệ.")

        records.sort(key=lambda item: item["stt"])

        # Xóa và nhập mới trong cùng một giao dịch.
        # Nếu có lỗi, SQLite sẽ rollback để giữ danh sách cũ.
        with connect() as conn:
            conn.execute("DELETE FROM participants")

            for person in records:
                conn.execute(
                    """
                    INSERT INTO participants
                        (code, full_name, role, organization, position)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        make_code(),
                        person["name"],
                        "Thành viên",
                        "",
                        person["position"],
                    ),
                )

        return RedirectResponse(
            "/admin?message="
            + quote(
                f"Đã thay thế danh sách cũ bằng danh sách mới gồm {len(records)} đại biểu."
            ),
            status_code=303,
        )

    except Exception as exc:
        return RedirectResponse(
            "/admin?error=" + quote(str(exc)),
            status_code=303,
        )

    finally:
        if wb is not None:
            wb.close()

        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            pass

        await file.close()


# =========================================================
# TẢI LÊN VĂN BẢN PDF
# =========================================================

@app.post("/admin/documents")
async def upload_document(
    request: Request,
    title: str = Form(""),
    file: UploadFile = File(...),
):
    """Chỉ admin mới được tải lên văn bản."""

    if not is_admin(request):
        return admin_login_redirect()

    if (
        not file.filename
        or Path(file.filename).suffix.lower() != ".pdf"
    ):
        return RedirectResponse(
            "/admin?error="
            + quote("Chỉ chấp nhận file PDF."),
            status_code=303,
        )

    stored_name = f"{uuid.uuid4().hex}.pdf"
    destination = UPLOAD_DIR / stored_name

    try:
        content = await file.read()

        if len(content) > 30 * 1024 * 1024:
            raise ValueError(
                "File PDF vượt quá giới hạn 30 MB."
            )

        if not content.startswith(b"%PDF-"):
            raise ValueError(
                "File không có định dạng PDF hợp lệ."
            )

        destination.write_bytes(content)

        clean_title = (
            normalized(title)
            or Path(file.filename).stem
        )

        with connect() as conn:
            conn.execute("""
                INSERT INTO documents (
                    title,
                    original_name,
                    stored_name
                )
                VALUES (?, ?, ?)
            """, (
                clean_title,
                Path(file.filename).name,
                stored_name,
            ))

        return RedirectResponse(
            "/admin?message="
            + quote("Đã tải lên văn bản."),
            status_code=303,
        )

    except Exception as exc:
        try:
            destination.unlink(missing_ok=True)
        except OSError:
            pass

        return RedirectResponse(
            "/admin?error=" + quote(str(exc)),
            status_code=303,
        )

    finally:
        await file.close()


# =========================================================
# ĐIỂM DANH ĐẠI BIỂU
# =========================================================

@app.post("/admin/attendance/{person_id}")
def toggle_attendance(
    request: Request,
    person_id: int,
    action: str = Form(...),
):
    """Đánh dấu có mặt hoặc hủy điểm danh."""

    if not is_admin(request):
        return admin_login_redirect()

    if action not in {"checkin", "undo"}:
        raise HTTPException(
            status_code=400,
            detail="Thao tác điểm danh không hợp lệ.",
        )

    with connect() as conn:
        person = conn.execute("""
            SELECT id
            FROM participants
            WHERE id = ?
        """, (person_id,)).fetchone()

        if not person:
            raise HTTPException(
                status_code=404,
                detail="Không tìm thấy đại biểu.",
            )

        if action == "checkin":
            conn.execute("""
                UPDATE participants
                SET
                    checked_in = 1,
                    checked_in_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (person_id,))

        else:
            conn.execute("""
                UPDATE participants
                SET
                    checked_in = 0,
                    checked_in_at = NULL
                WHERE id = ?
            """, (person_id,))

    return RedirectResponse(
        "/admin",
        status_code=303,
    )


# =========================================================
# XÓA VĂN BẢN
# =========================================================

@app.post("/admin/documents/{doc_id}/delete")
def delete_document(
    request: Request,
    doc_id: int,
):
    """Chỉ admin mới được xóa văn bản."""

    if not is_admin(request):
        return admin_login_redirect()

    with connect() as conn:
        document = conn.execute("""
            SELECT *
            FROM documents
            WHERE id = ?
        """, (doc_id,)).fetchone()

        if not document:
            raise HTTPException(
                status_code=404,
                detail="Không tìm thấy văn bản.",
            )

        conn.execute("""
            DELETE FROM documents
            WHERE id = ?
        """, (doc_id,))

    # Xóa file PDF khỏi thư mục lưu trữ.
    try:
        (
            UPLOAD_DIR / document["stored_name"]
        ).unlink(missing_ok=True)

    except OSError:
        # Nếu không xóa được file vật lý,
        # bản ghi database vẫn đã được xóa.
        pass

    return RedirectResponse(
        "/admin?message=" + quote("Đã xóa văn bản."),
        status_code=303,
    )


# =========================================================
# XEM PDF CÔNG KHAI
# =========================================================

@app.get("/documents/{doc_id}/view")
def view_document(doc_id: int):
    """
    Cho phép mọi người xem PDF mà không cần đăng nhập.
    """

    with connect() as conn:
        document = conn.execute("""
            SELECT *
            FROM documents
            WHERE id = ?
        """, (doc_id,)).fetchone()

    if not document:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy văn bản.",
        )

    file_path = UPLOAD_DIR / document["stored_name"]

    if not file_path.is_file():
        raise HTTPException(
            status_code=404,
            detail="Tệp PDF không tồn tại trên máy chủ.",
        )

    return FileResponse(
        path=file_path,
        media_type="application/pdf",
        filename=document["original_name"],
        content_disposition_type="inline",
    )


# =========================================================
# HEALTH CHECK CHO RENDER
# =========================================================

@app.get("/health")
def health():
    return {
        "status": "ok"
    }